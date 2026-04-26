"""Tests for Excel exporter, including row-level data_quality_warning surfacing."""

from datetime import date, time
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from src.cascade.cascade_detector import compute_kpis, detect_cascade
from src.export.excel_exporter import export_to_excel
from src.models.event import AirportClosureEvent


def _han_event() -> AirportClosureEvent:
    return AirportClosureEvent(
        airport="HAN",
        closure_date=date(2026, 4, 24),
        start_time=time(14, 0),
        end_time=time(18, 0),
    )


def _sample_df() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "flight_no": "VN123",
                "aircraft_reg": "VN-A500",
                "aircraft_type": "A321",
                "origin": "SGN",
                "destination": "HAN",
                "std": time(13, 0),
                "sta": time(15, 0),
                "data_quality_warning": "TIME_PARSE_FAILED at sta",
            },
            {
                "flight_no": "VN124",
                "aircraft_reg": "VN-A500",
                "aircraft_type": "A321",
                "origin": "HAN",
                "destination": "DAD",
                "std": time(16, 0),
                "sta": time(17, 30),
                "data_quality_warning": None,
            },
            {
                "flight_no": "VN200",
                "aircraft_reg": "VN-B100",
                "aircraft_type": "A320",
                "origin": "SGN",
                "destination": "DAD",
                "std": time(8, 0),
                "sta": time(9, 30),
                "data_quality_warning": None,
            },
        ]
    )


def _read_sheet(buf: BytesIO, sheet: str) -> pd.DataFrame:
    buf.seek(0)
    return pd.read_excel(buf, sheet_name=sheet)


def test_export_returns_bytes_io():
    df = detect_cascade(_sample_df(), _han_event())
    kpis = compute_kpis(df)

    out = export_to_excel(df, _han_event(), kpis, warnings=["W1", "W2"])

    assert isinstance(out, BytesIO)
    out.seek(0)
    wb = load_workbook(out, read_only=True)
    expected_sheets = {
        "Parameters",
        "KPI Summary",
        "Affected Flights",
        "All Flights",
        "Data Quality Warnings",
    }
    assert expected_sheets.issubset(set(wb.sheetnames))


def test_affected_sheet_includes_data_quality_warning_column():
    df = detect_cascade(_sample_df(), _han_event())
    kpis = compute_kpis(df)

    out = export_to_excel(df, _han_event(), kpis, warnings=[])
    affected = _read_sheet(out, "Affected Flights")

    assert "data_quality_warning" in affected.columns
    # Row-level warning must reach the Excel output
    matching = affected[affected["flight_no"] == "VN123"]
    assert not matching.empty
    assert matching.iloc[0]["data_quality_warning"] == "TIME_PARSE_FAILED at sta"


def test_warnings_sheet_handles_empty_list():
    df = detect_cascade(_sample_df(), _han_event())
    kpis = compute_kpis(df)

    out = export_to_excel(df, _han_event(), kpis, warnings=[])
    warns = _read_sheet(out, "Data Quality Warnings")

    assert "Warning" in warns.columns
    assert warns["Warning"].tolist() == ["No warnings"]
